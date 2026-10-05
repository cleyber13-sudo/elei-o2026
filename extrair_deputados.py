#!/usr/bin/env python3
"""Extrai deputados federais, estaduais e distritais eleitos em 2026 e os
3 primeiros suplentes (por partido/federação, em cada UF), com nome completo,
partido, UF, CPF e cargo.

Fontes (TSE):
  - resultados.tse.jus.br         -> candidatos, votação e situação de totalização
      oficial/comum/config/ele-c.json
      oficial/ele2026/{cd}/dados/{uf}/{uf}-c{cargo:04d}-e{cd:06d}-u.json
  - divulgacandcontas.tse.jus.br  -> ficha do candidato (CPF)
      divulga/rest/v1/candidatura/buscar/2026/{UF}/{eleicao}/candidato/{sqcand}
  - opcional: CSV consulta_cand_2026 dos dados abertos do TSE (--cpf-csv), usado
    quando o DivulgaCandContas não estiver acessível.

Uso:
  python3 extrair_deputados.py                 # todas as UFs
  python3 extrair_deputados.py --ufs CE SP     # apenas algumas UFs
  python3 extrair_deputados.py --cpf-csv consulta_cand_2026_BRASIL.csv
  python3 extrair_deputados.py --xlsx deputados_2026.xlsx   # requer openpyxl

Só usa a biblioteca padrão do Python 3.8+ (openpyxl apenas para --xlsx).
"""
import argparse
import csv
import json
import sys
import time
import urllib.error
import urllib.request
from concurrent.futures import ThreadPoolExecutor

ANO = 2026
DIVULGA = "https://divulgacandcontas.tse.jus.br/divulga/rest/v1"
RESULTADOS = "https://resultados.tse.jus.br/oficial"
# Id da eleição geral 2026 no DivulgaCandContas (da URL da ficha do candidato).
ELEICAO_DIVULGA_PADRAO = "20322002026"

UFS = ["AC", "AL", "AM", "AP", "BA", "CE", "DF", "ES", "GO", "MA", "MG", "MS",
       "MT", "PA", "PB", "PE", "PI", "PR", "RJ", "RN", "RO", "RR", "RS", "SC",
       "SE", "SP", "TO"]
DEP_FEDERAL, DEP_ESTADUAL, DEP_DISTRITAL = 6, 7, 8
UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/129.0 Safari/537.36")


def get_json(url, tentativas=4):
    """JSON da URL; None em 404. Erros 403 (bloqueio) não são repetidos."""
    for i in range(tentativas):
        try:
            req = urllib.request.Request(url, headers={"User-Agent": UA, "Accept": "application/json"})
            with urllib.request.urlopen(req, timeout=60) as r:
                return json.loads(r.read().decode("utf-8"))
        except urllib.error.HTTPError as e:
            if e.code == 404:
                return None
            if e.code == 403 or i == tentativas - 1:
                raise
        except (urllib.error.URLError, TimeoutError):
            if i == tentativas - 1:
                raise
        time.sleep(2 ** (i + 1))


def eleicao_resultados(cargo):
    """Código (cd) da eleição de 1º turno de 2026 que contém o cargo."""
    cfg = get_json(f"{RESULTADOS}/comum/config/ele-c.json") or {}
    for pleito in cfg.get("pl", []):
        if pleito.get("c") != f"ele{ANO}":
            continue
        for e in pleito.get("e", []):
            cargos = {c.get("cd") for a in e.get("abr", []) for c in a.get("cp", [])}
            if str(e.get("t")) == "1" and str(cargo) in cargos:
                return str(e["cd"])
    raise SystemExit(f"cargo {cargo} não encontrado em ele-c.json para {ANO}")


def votacao(uf, cargo, cd):
    """Arquivo de resultado do cargo na UF: candidatos agrupados por agremiação."""
    url = (f"{RESULTADOS}/ele{ANO}/{cd}/dados/{uf.lower()}/"
           f"{uf.lower()}-c{cargo:04d}-e{int(cd):06d}-u.json")
    dados = get_json(url) or {}
    carg = (dados.get("carg") or [{}])[0]
    candidatos = []
    for agr in carg.get("agr", []):
        for par in agr.get("par", []):
            for c in par.get("cand", []):
                candidatos.append({"c": c, "par": par, "agr": agr})
    info = {
        "vagas": int(carg.get("nv") or 0),
        "cargo": carg.get("nmn", ""),
        "totalizacao_final": dados.get("tf") == "s",
        "secoes_pct": (dados.get("s") or {}).get("pst", ""),
        "atualizado": f"{dados.get('dt', '')} {dados.get('ht', '')}".strip(),
    }
    return candidatos, info


def votos(c):
    try:
        return int(str(c.get("vap", "0")).replace(".", ""))
    except ValueError:
        return 0


def selecionar(uf, cargo, cd, n_suplentes):
    candidatos, info = votacao(uf, cargo, cd)
    eleitos, suplentes = [], {}
    for it in candidatos:
        s = (it["c"].get("st") or "").lower()
        it["grupo"] = it["agr"].get("nm") or it["par"].get("sg", "")
        if s.startswith("eleito"):
            eleitos.append(it)
        elif s.startswith("suplente"):
            suplentes.setdefault(it["grupo"], []).append(it)
    selecionados = [(it, "Eleito", None) for it in sorted(eleitos, key=lambda it: -votos(it["c"]))]
    for grupo in sorted(suplentes):
        lista = sorted(suplentes[grupo], key=lambda it: -votos(it["c"]))
        for pos, it in enumerate(lista[:n_suplentes], 1):
            selecionados.append((it, "Suplente", pos))
    return selecionados, len(candidatos), info


def carregar_cpf_csv(caminho):
    """{SQ_CANDIDATO: CPF} a partir do consulta_cand_2026 (dados abertos do TSE)."""
    cpfs = {}
    with open(caminho, encoding="latin-1", newline="") as fh:
        for linha in csv.DictReader(fh, delimiter=";"):
            sq, cpf = linha.get("SQ_CANDIDATO"), linha.get("NR_CPF_CANDIDATO")
            if sq and cpf and cpf not in ("-1", "-4"):
                cpfs[sq.strip()] = cpf.strip().zfill(11)
    return cpfs


def divulga_acessivel(uf, eleicao, sqcand):
    try:
        get_json(f"{DIVULGA}/candidatura/buscar/{ANO}/{uf}/{eleicao}/candidato/{sqcand}", tentativas=2)
        return True
    except Exception as exc:  # noqa: BLE001
        print(f"aviso: DivulgaCandContas inacessível ({exc}); CPF só via --cpf-csv", file=sys.stderr)
        return False


def ficha(uf, eleicao, sqcand):
    try:
        return get_json(f"{DIVULGA}/candidatura/buscar/{ANO}/{uf}/{eleicao}/candidato/{sqcand}") or {}
    except Exception:  # noqa: BLE001
        return {}


def formatar_cpf(cpf):
    d = "".join(ch for ch in str(cpf or "") if ch.isdigit())
    return f"{d[:3]}.{d[3:6]}.{d[6:9]}-{d[9:]}" if len(d) == 11 else ""


def gravar_xlsx(caminho, linhas, colunas):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    abas = {"Deputado Federal": "Dep. Federal", "Deputado Estadual": "Dep. Estadual",
            "Deputado Distrital": "Dep. Distrital"}
    for cargo, nome in abas.items():
        ws = wb.create_sheet(nome)
        ws.append(colunas)
        for c in ws[1]:
            c.font = Font(bold=True)
        for l in linhas:
            if l["cargo"] == cargo:
                ws.append([l[k] for k in colunas])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, k in enumerate(colunas, 1):
            largura = max([len(str(k))] + [len(str(l[k])) for l in linhas if l["cargo"] == cargo])
            ws.column_dimensions[get_column_letter(i)].width = min(largura + 2, 50)
    wb.save(caminho)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--ufs", nargs="*", default=UFS)
    ap.add_argument("--suplentes", type=int, default=3)
    ap.add_argument("--eleicao", default=ELEICAO_DIVULGA_PADRAO, help="id da eleição no DivulgaCandContas")
    ap.add_argument("--cpf-csv", help="consulta_cand_2026 (CSV ';' latin-1) com SQ_CANDIDATO e NR_CPF_CANDIDATO")
    ap.add_argument("--saida", default=f"deputados_eleitos_suplentes_{ANO}.csv")
    ap.add_argument("--xlsx", help="também grava um .xlsx com uma aba por cargo (requer openpyxl)")
    ap.add_argument("--threads", type=int, default=8)
    args = ap.parse_args()

    tarefas, resumo = [], []
    for uf in [u.upper() for u in args.ufs]:
        for cargo in [DEP_FEDERAL, DEP_DISTRITAL if uf == "DF" else DEP_ESTADUAL]:
            cd = eleicao_resultados(cargo)
            sel, total, info = selecionar(uf, cargo, cd, args.suplentes)
            n_el = sum(1 for _, t, _ in sel if t == "Eleito")
            resumo.append((uf, cargo, n_el, info))
            print(f"{uf} cargo {cargo}: {total} candidatos, {n_el}/{info['vagas']} eleitos, "
                  f"{len(sel) - n_el} suplentes (seções {info['secoes_pct']}%, "
                  f"{'final' if info['totalizacao_final'] else 'PARCIAL'}, {info['atualizado'] or 'n/d'})",
                  file=sys.stderr)
            tarefas += [(uf, info["cargo"], it, tipo, pos) for it, tipo, pos in sel]

    cpfs = carregar_cpf_csv(args.cpf_csv) if args.cpf_csv else {}
    usar_divulga = bool(tarefas) and divulga_acessivel(tarefas[0][0], args.eleicao, tarefas[0][2]["c"]["sqcand"])

    def montar(t):
        uf, cargo, it, tipo, pos = t
        c, par = it["c"], it["par"]
        sq = str(c.get("sqcand"))
        cpf = cpfs.get(sq, "")
        nome = c.get("nm", "")
        if usar_divulga and not cpf:
            f = ficha(uf, args.eleicao, sq)
            cpf, nome = f.get("cpf") or "", f.get("nomeCompleto") or nome
        return {
            "uf": uf,
            "cargo": cargo,
            "resultado": tipo,
            "ordem_suplencia": pos or "",
            "situacao_totalizacao": c.get("st", ""),
            "nome_completo": nome,
            "nome_urna": c.get("nmu", ""),
            "numero": c.get("n", ""),
            "partido_sigla": par.get("sg", ""),
            "partido_nome": par.get("nm", ""),
            "partido_federacao": it["grupo"],
            "cpf": formatar_cpf(cpf),
            "votos": votos(c),
            "id_candidato_tse": sq,
        }

    with ThreadPoolExecutor(args.threads) as ex:
        linhas = list(ex.map(montar, tarefas))

    colunas = list(linhas[0].keys()) if linhas else ["uf"]
    with open(args.saida, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=colunas, delimiter=";")
        w.writeheader()
        w.writerows(linhas)
    print(f"{len(linhas)} linhas gravadas em {args.saida}", file=sys.stderr)
    if args.xlsx:
        gravar_xlsx(args.xlsx, linhas, colunas)
        print(f"planilha gravada em {args.xlsx}", file=sys.stderr)

    fed = sum(n for _, cg, n, _ in resumo if cg == DEP_FEDERAL)
    est = sum(n for _, cg, n, _ in resumo if cg != DEP_FEDERAL)
    print(f"eleitos: {fed} federais, {est} estaduais/distritais", file=sys.stderr)
    for uf, cg, n, info in resumo:
        if not info["totalizacao_final"] or n != info["vagas"]:
            print(f"ATENÇÃO {uf} cargo {cg}: {n}/{info['vagas']} eleitos, totalização "
                  f"{'final' if info['totalizacao_final'] else 'parcial'}", file=sys.stderr)
    sem_cpf = sum(1 for l in linhas if not l["cpf"])
    if sem_cpf:
        print(f"ATENÇÃO: {sem_cpf} linhas sem CPF", file=sys.stderr)


if __name__ == "__main__":
    main()
